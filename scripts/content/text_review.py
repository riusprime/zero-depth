#!/usr/bin/env python3
"""v0.8.0 (owner, 2026-10-17): every piece of on-screen text in one workbook for the owner to review and rename.

Reads locale/strings.csv (every visible string goes through tr(), EI-10) and finds where each key is used in src/ and
data/, then names the screen it shows on. Keys built at run time (e.g. UI_OPT_<value>, CAUSE_<attack>) take the
screens of their siblings. Writes build/TEXT_REVIEW.xlsx; the lead copies it to docs/TEXT_REVIEW.xlsx.

Usage: python3 scripts/content/text_review.py
"""
import collections
import csv
import os
import re

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.worksheet.datavalidation import DataValidation

SCRIPT_SCREENS = {
    "main_menu": "Main menu",
    "pause_menu": "Pause menu",
    "options_menu": "Options",
    "option_cycler": "Options",
    "step_slider": "Options",
    "credits_view": "Credits",
    "codex_view": "Codex",
    "codex_tile": "Codex",
    "achievements_page": "Achievements page",
    "achievement_row": "Achievements page",
    "achievement_toast": "Achievement pop-up",
    "build_picker": "Weapon pick (new run)",
    "build_card": "Weapon pick (new run)",
    "build_page": "Build page (pause → Build)",
    "wand_row": "Build page (pause → Build)",
    "end_panel": "End of run screen",
    "input_labels": "Button prompts",
    "menu_panel": "Menus (shared)",
    "hud": "HUD (in play)",
    "hud_bar": "HUD: health bar",
    "boss_bar": "HUD: boss bar",
    "ability_hud": "HUD: ability slots",
    "build_hud": "HUD: build strip",
    "kit_hud": "HUD: kit",
    "danger_meter": "HUD: danger meter",
    "heat_meter": "HUD: heat meter",
    "dps_meter": "Playground: DPS meter",
    "phase_hud": "HUD: phase line",
    "event_hud": "HUD: floor events",
    "event_panel": "Floor event panel",
    "overrun_hud": "HUD: Overrun",
    "threat_panel": "Threat / curses panel",
    "minimap": "Minimap",
    "minimap_view": "Minimap",
    "onboarding_hints": "Tutorial hints",
    "pick_panel": "Card pick (altars, chests)",
    "pick_slot": "Card pick (altars, chests)",
    "item_card": "Card pick (altars, chests)",
    "card_gain": "Card pick: what gets better",
    "card_deal": "Card pick (altars, chests)",
    "stone_card": "Card pick (altars, chests)",
    "combo_card": "Combo cards",
    "swap_panel": "Swap panel (slots full)",
    "shop_panel": "Shop",
    "shop_hud": "Shop",
    "shop_tile": "Shop",
    "gamble_card": "Gamble shrine",
    "gamble_hud": "Gamble shrine",
    "gamble_stats_panel": "Gamble shrine",
    "modal_panel": "Pop-up panels",
    "loading_cover": "Loading screen",
    "plaques": "Card plaques",
    "gamble_icons": "Gamble shrine",
    "boss_causes": "End of run screen: what killed you",
    "main": "Game messages (in play)",
    "playground_session": "Playground (dev)",
    "minimap_state": "Minimap",
    "credits_definition": "Credits",
}
DATA_SCREENS = {
    "items": "Cards: names and descriptions",
    "achievements": "Achievements",
    "bosses": "Bosses: names and attacks",
    "stat_cards": "Stat cards",
    "curses": "Curses",
    "combos": "Combo cards",
    "events": "Floor events",
    "enemies": "Enemy names",
    "audio": "Captions (subtitles for sounds)",
    "abilities": "Abilities",
    "final_boss": "Final boss",
    "curves": "Danger tiers",
    "biomes": "Biome names",
    "builds": "Weapon pick (new run)",
    "utilities": "Utility cards",
    "run": "Run",
    "wand": "Build page (pause → Build)",
}
PREFIX_SCREENS = {
    "CAUSE": "End of run screen: what killed you",
    "CAPTION": "Captions (subtitles for sounds)",
    "ACH": "Achievements",
    "ITEM": "Cards: names and descriptions",
    "HINT": "Tutorial hints",
    "BIOME": "Biome names",
    "ENEMY": "Enemy names",
}
DEV = "Dev tools (hidden from players)"
## Code that only passes a key along (readers, looks, events): it names no screen.
PLUMBING = {
    "world_reader", "item_looks", "item_visuals", "attack_view", "attack_form_looks", "audio_events", "sim_event",
    "core_views", "player_avatar", "skill_visuals", "world_view_root", "portal_transit_view", "shop_stall_view",
    "shop_terminal_view", "playground_api",
}


def screen_of(path: str) -> str:
    parts = path.split(os.sep)
    if parts[0] == "src" and len(parts) > 1 and parts[1] == "debug" or "dev_" in os.path.basename(path):
        return DEV
    if parts[0] == "data" and len(parts) > 1:
        return DATA_SCREENS.get(parts[1], parts[1].replace("_", " ").capitalize())
    base = os.path.splitext(os.path.basename(path))[0]
    if base in SCRIPT_SCREENS:
        return SCRIPT_SCREENS[base]
    return base.replace("_", " ").capitalize()


def main() -> None:
    rows = list(csv.reader(open("locale/strings.csv", encoding="utf-8")))
    header, body = rows[0], rows[1:]
    i_en, i_es = header.index("en"), header.index("es")
    files = []
    for root in ("src", "data"):
        for d, _, fs in os.walk(root):
            files += [os.path.join(d, f) for f in fs if f.endswith((".gd", ".tres", ".tscn"))]
    text = {f: open(f, encoding="utf-8", errors="ignore").read() for f in files}
    used = collections.defaultdict(list)
    for r in body:
        pat = re.compile(r"(?<![A-Z0-9_])" + re.escape(r[0]) + r"(?![A-Z0-9_])")
        used[r[0]] = sorted(f for f, t in text.items() if pat.search(t))
    # Keys built at run time: borrow the files of the siblings sharing the longest prefix.
    found = [k for k in used if used[k]]
    for k in [k for k in used if not used[k]]:
        best, best_len = None, 0
        for o in found:
            n = len(os.path.commonprefix([k, o]))
            if n > best_len:
                best, best_len = o, n
        if best and best_len >= len(k.split("_")[0]) + 2:
            used[k] = used[best]
    out = []
    for r in body:
        k = r[0]
        real = [f for f in used[k] if os.path.splitext(os.path.basename(f))[0] not in PLUMBING]
        real = [f for f in real if not f.startswith(("src/sim", "src/application", "src/content"))] or real
        screens = sorted({screen_of(f) for f in real})
        player = [s for s in screens if s != DEV]
        if not player and k.split("_")[0] in PREFIX_SCREENS and screens != [DEV]:
            player = [PREFIX_SCREENS[k.split("_")[0]]]
        where = ", ".join(player) if player else (DEV if screens else "Not found in code (check)")
        out.append([where, k, r[i_en], r[i_es], ", ".join(used[k][:4]) + (" …" if len(used[k]) > 4 else "")])
    out.sort(key=lambda x: (x[0] == DEV, x[0] == "Not found in code (check)", x[0], x[1]))

    wb = Workbook()
    leg = wb.active
    leg.title = "How to use"
    lines = [
        ("Text review: every string the game shows", True),
        ("", False),
        ("1. Go to the 'All text' sheet. Each row is one piece of text; use the filter on 'Where it appears' to see one screen.", False),
        ("2. To change a text, type the new wording in the yellow columns 'New English' and/or 'New Spanish'. Leave them empty to keep the current text.", False),
        ("3. 'Status': pick OK, Change or Remove (optional). Add anything else in 'Notes'.", False),
        ("4. Send the file back; the lead applies every yellow cell to locale/strings.csv and re-checks the screens.", False),
        ("", False),
        ("Placeholders like {0}, %d or %s are filled in by the game (numbers, names): keep them in your new text.", False),
        ("Zero Depth is the game's name; keep it as written in both languages.", False),
        ("Rows marked 'Dev tools (hidden from players)' are at the bottom; players never see them.", False),
        ("", False),
        ("Example row (how a change looks):", True),
        ("Where it appears: Main menu · Key: UI_PLAY · English: Play · Spanish: Jugar · New English: Start · New Spanish: Empezar · Status: Change", False),
        ("", False),
        (f"Source: locale/strings.csv ({len(out)} strings), generated by scripts/content/text_review.py.", False),
    ]
    for n, (t, b) in enumerate(lines, 1):
        c = leg.cell(row=n, column=1, value=t)
        c.font = Font(name="Arial", bold=b, size=12 if b and n == 1 else 10)
    leg.column_dimensions["A"].width = 130

    ws = wb.create_sheet("All text")
    cols = ["Where it appears", "Key", "English (now)", "Spanish (now)", "New English", "New Spanish", "Status", "Notes", "Code files"]
    widths = [32, 30, 48, 48, 40, 40, 11, 30, 50]
    ws.append(cols)
    for j, w in enumerate(widths, 1):
        ws.column_dimensions[ws.cell(row=1, column=j).column_letter].width = w
    hf = Font(name="Arial", bold=True, color="FFFFFF")
    for c in ws[1]:
        c.font = hf
        c.fill = PatternFill("solid", fgColor="3A3A3E")
        c.alignment = Alignment(vertical="center", wrap_text=True)
    yellow = PatternFill("solid", fgColor="FFF2B3")
    grey = PatternFill("solid", fgColor="EEEEEE")
    body_font = Font(name="Arial", size=10)
    dim_font = Font(name="Arial", size=9, color="777777")
    for r in out:
        ws.append([r[0], r[1], r[2], r[3], None, None, None, None, r[4]])
        n = ws.max_row
        for j in range(1, 10):
            c = ws.cell(row=n, column=j)
            c.font = dim_font if j in (2, 9) else body_font
            c.alignment = Alignment(vertical="top", wrap_text=j in (3, 4, 5, 6, 8))
            if j in (5, 6):
                c.fill = yellow
            if r[0] == DEV and j in (1, 3, 4):
                c.fill = grey
    dv = DataValidation(type="list", formula1='"OK,Change,Remove"', allow_blank=True)
    ws.add_data_validation(dv)
    dv.add(f"G2:G{ws.max_row}")
    ws.freeze_panes = "C2"
    ws.auto_filter.ref = f"A1:I{ws.max_row}"

    summ = wb.create_sheet("Screens")
    summ.append(["Screen", "Strings"])
    for c in summ[1]:
        c.font = hf
        c.fill = PatternFill("solid", fgColor="3A3A3E")
    counts = collections.Counter(r[0] for r in out)
    for s, n in sorted(counts.items(), key=lambda x: (x[0] == DEV, -x[1])):
        summ.append([s, n])
    for row in summ.iter_rows(min_row=2):
        for c in row:
            c.font = body_font
    summ.column_dimensions["A"].width = 60
    summ.column_dimensions["B"].width = 10
    os.makedirs("build", exist_ok=True)
    wb.save("build/TEXT_REVIEW.xlsx")
    print(f"wrote build/TEXT_REVIEW.xlsx ({len(out)} strings, {len(counts)} screens)")
    for s, n in counts.most_common():
        print(f"  {n:4d}  {s}")


if __name__ == "__main__":
    main()
